from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path('docs/operations/us-common-stocks-100pct-risers-20261015.docx')
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)

styles = doc.styles
for stylename in ['Normal','Title','Heading 1','Heading 2']:
    s = styles[stylename]
    s.font.name = 'Aptos'
    s._element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')
styles['Normal'].font.size = Pt(9.5)
styles['Normal'].paragraph_format.space_after = Pt(5)
styles['Title'].font.size = Pt(23); styles['Title'].font.bold = True
styles['Title'].font.color.rgb = RGBColor(20, 45, 74)
styles['Heading 1'].font.size = Pt(15); styles['Heading 1'].font.color.rgb = RGBColor(20, 65, 105)
styles['Heading 2'].font.size = Pt(11); styles['Heading 2'].font.color.rgb = RGBColor(34, 86, 125)

def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); pr.append(shd)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr(); el = OxmlElement('w:tblHeader'); el.set(qn('w:val'), 'true'); trPr.append(el)

def link(p, label, url):
    rid = p.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), rid)
    r = OxmlElement('w:r'); rp = OxmlElement('w:rPr'); color = OxmlElement('w:color'); color.set(qn('w:val'),'0563C1'); rp.append(color)
    under = OxmlElement('w:u'); under.set(qn('w:val'),'single'); rp.append(under); r.append(rp)
    t = OxmlElement('w:t'); t.text = label; r.append(t); h.append(r); p._p.append(h)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = p.add_run('미국 상장 보통주 | 2026년 10월 15일 기준'); r.bold=True; r.font.color.rgb=RGBColor(90,105,120); r.font.size=Pt(9)
doc.add_heading('최근 30일 장중 고점 기준 100% 이상 상승 종목', 0)
p=doc.add_paragraph('미국 보통주 · 시가총액 5억 달러 이하')
p.runs[0].font.size=Pt(12); p.runs[0].font.color.rgb=RGBColor(75,90,105)

doc.add_heading('요약', 1)
doc.add_paragraph('2026년 9월 14일 장중 고가를 기준점으로 두고, 그날부터 10월 14일까지 30일 구간에 기록된 최고 장중가까지의 상승률로 선별했습니다. 확인된 조건 충족 종목은 PDS Biotechnology (PDSB), TJGC Group (TJGC), Nautilus Biotechnology (NAUT)입니다. 세 종목 모두 미국 거래소 상장 보통주이며, 2026년 10월 14일 기준 시가총액이 5억 달러 미만입니다.')

table=doc.add_table(rows=1, cols=6); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Light Shading Accent 1'
headers=['티커 / 회사','증권 종류','9/14 장중 고가','30일 구간 최고가','고점 기준 상승률','시가총액']
for c,t in zip(table.rows[0].cells,headers): c.text=t; shade(c,'D9EAF4')
set_repeat_table_header(table.rows[0])
rows=[
 ['PDSB\nPDS Biotechnology','NASDAQ 보통주','$0.45','$1.79 (10/13)','+297.8%','$102.19M'],
 ['TJGC\nTJGC Group','NASDAQ Class A 보통주','$12.38','$36.76 (9/29)','+196.9%','$324.31M'],
 ['NAUT\nNautilus Biotechnology','NASDAQ 보통주','$1.02','$2.475 (10/5)','+142.6%','$274.87M'],
]
for row in rows:
    cells=table.add_row().cells
    for c,t in zip(cells,row): c.text=t; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

doc.add_paragraph('수익률 계산: (2026년 9월 14일 장중 고가부터 10월 14일까지의 최고 장중가 ÷ 9월 14일 장중 고가 − 1) × 100. 계산값은 반올림했습니다. 이는 고점 대 고점의 최대 구간 상승폭이며, 실제 투자자가 그 저점·고점에 체결했음을 뜻하지 않습니다.')

doc.add_heading('종목별 확인 사항',1)
doc.add_heading('PDSB | PDS Biotechnology',2)
doc.add_paragraph('가격·성과: 9월 14일 장중 고가 $0.45에서 10월 13일 장중 고가 $1.79까지 상승해 약 +297.8%입니다. 10월 14일 시가총액은 $102.19M입니다. 9월 8일 회사 발표는 Nant Capital 주도의 최대 $22.55M PIPE를 알렸고, 초기 자금 조달과 조건부 후속 조달, 보통주 및 워런트 발행 구조를 설명했습니다. 이는 시장이 주목할 수 있는 촉매인 동시에 잠재적 주식 희석 요인입니다.')
doc.add_paragraph('주의 위험: 해당 PIPE에는 보통주·사전납입 워런트·커먼 워런트가 포함됩니다. 급격한 단기 상승과 높은 거래량은 가격 변동성을 키울 수 있습니다. 임상 결과와 자금 조달은 투자 성과를 보장하지 않습니다.')

doc.add_heading('TJGC | TJGC Group',2)
doc.add_paragraph('가격·성과: 9월 14일 장중 고가 $12.38에서 30일 구간 중 9월 29일 고가 $36.76까지 상승해 약 +196.9%입니다. 10월 14일 시가총액은 $324.31M입니다. 회사의 SEC Form 6-K에 따르면 9월 24일 이사회는 최대 $2.0M의 Class A ordinary shares 자사주 매입 프로그램을 승인했습니다. 이 프로그램은 재량적이며 회사가 실제 매입을 해야 할 의무는 없습니다.')
doc.add_paragraph('주의 위험: 거래량은 일부 세션에서 매우 낮았습니다(예: 10월 14일 약 68,904주). 낮은 유동성은 호가 간격 확대와 큰 가격 충격을 유발할 수 있습니다. 30일 상승률이 크게 높더라도 현재 가격에서의 매수 가능성이나 지속성을 뜻하지 않습니다.')

doc.add_heading('NAUT | Nautilus Biotechnology',2)
doc.add_paragraph('가격·성과: 9월 14일 장중 고가 $1.02에서 30일 구간 중 10월 5일 장중 고가 $2.475까지 상승해 약 +142.6%입니다. 10월 14일 시가총액은 $274.87M입니다. SEC 공시에 따르면 회사는 9월 11일 최대 $125M 규모의 ATM 판매 계약을 맺어 필요에 따라 보통주를 발행·매각할 수 있습니다.')
doc.add_paragraph('주의 위험: $125M ATM은 실제 매각 여부·규모가 회사 재량이지만, 실행 시 주식 공급과 기존 주주 희석 가능성이 있습니다. 특히 30일 상승률이 급격하고 시장 데이터 제공처 간 종가도 소폭 다르므로 체결 시점과 유동성을 확인해야 합니다.')

doc.add_heading('기준 및 해석',1)
for txt in [
 '시장·종목: 미국 거래소 상장 보통주만 포함했습니다. ETF, ETN, 우선주, 워런트, 유닛 및 인수목적회사 증권은 제외했습니다. TJGC는 SEC 공시상 Nasdaq 상장 Class A ordinary shares로 확인했습니다.',
 '기간 및 고점 규칙: 2026년 9월 14일부터 10월 14일까지 미국 동부시간 거래일의 장중 가격을 사용했습니다. 출발점은 9월 14일 고가이며, 구간 최고가는 9월 14일을 포함해 10월 14일까지 각 거래일의 고가 중 최댓값입니다. 따라서 측정값은 고점 대 고점이며, 종가 수익률이나 30일 최저점부터의 상승률이 아닙니다.',
 '시가총액: 2026년 10월 14일 종가 기준으로 제시된 시장 데이터 제공처의 값입니다. 시가총액은 주가와 발행주식수에 따라 바뀌며, 데이터 제공처별 집계·시점 차이가 있을 수 있습니다.',
 '범위: 이 문서는 가격·시총 필터 결과이지 매수 추천이나 전체 시장에 대한 공식적 완전성 보증이 아닙니다. 극단적 상승 종목은 변동성·유동성·희석 위험을 별도로 검토해야 합니다.'
]: doc.add_paragraph(txt,style='List Bullet')

doc.add_heading('출처',1)
sources=[
 ('PDSB 가격 이력 및 장중 고가','https://stockanalysis.com/stocks/pdsb/history/'),
 ('PDSB 2026-10-14 시가총액','https://stockanalysis.com/stocks/pdsb/market-cap/'),
 ('PDS Biotechnology 2026-09-08 자금 조달 발표','https://www.globenewswire.com/news-release/2026/09/08/3357657/0/en/pds-biotech-announces-up-to-22-55-million-financing-led-by-dr-patrick-soon-shiong-founder-of-nantworks.html'),
 ('TJGC 가격 이력 및 장중 고가','https://chartexchange.com/symbol/nasdaq-tjgc/historical/'),
 ('TJGC 2026-10-14 시가총액','https://stockanalysis.com/stocks/tjgc/market-cap/'),
 ('TJGC SEC Form 6-K: 자사주 매입 프로그램','https://www.sec.gov/Archives/edgar/data/1969928/000118518526004251/tjgc6k092426.htm'),
 ('TJGC SEC 등록 공시: Nasdaq 상장 Class A Ordinary Shares','https://www.sec.gov/Archives/edgar/data/1969928/000118518526003808/tjgcf3081926.htm'),
 ('NAUT 가격 이력 및 장중 고가','https://chartexchange.com/symbol/nasdaq-naut/historical/'),
 ('NAUT 2026-10-14 시가총액','https://stockanalysis.com/stocks/naut/market-cap/'),
 ('NAUT SEC Form 8-K: 보통주 및 ATM 계약','https://www.sec.gov/Archives/edgar/data/1808805/000180880526000041/naut-20260911.htm'),
]
for title,url in sources:
    p=doc.add_paragraph(style='List Bullet'); link(p,title,url)

doc.add_paragraph('시장 정보 기준 시각: 2026년 10월 14일 종가. 문서 작성일: 2026년 10월 15일.')
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT.resolve())
